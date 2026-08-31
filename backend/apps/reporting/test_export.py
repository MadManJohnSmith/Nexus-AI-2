"""
Pruebas unitarias y de integración para el motor de exportación tabular y documental (HU-28).
"""
import io
import openpyxl
from django.utils import timezone
from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from rest_framework import status

from apps.students.models import Student, Semester, AcademicCommittee
from apps.tutoring.models import TutoringSession, TutoringParticipant, TutoringObservation
from apps.agreements.models import Agreement, AgreementAuditLog
from apps.thesis.models import ThesisProgress
from apps.academic_output.models import Publication, AcademicEvent, ResearchStay, OtherProduct
from apps.evidence.models import Evidence

User = get_user_model()


class ExportEngineAPITests(APITestCase):
    def setUp(self):
        # 1. Crear usuarios con distintos roles
        self.coordinator = User.objects.create_user(
            email='coord@test.edu',
            password='password123',
            role='COORDINADOR',
            first_name='Laura',
            last_name='Coordinadora'
        )

        self.advisor_assigned = User.objects.create_user(
            email='advisor1@test.edu',
            password='password123',
            role='ASESOR',
            first_name='Carlos',
            last_name='Asesor'
        )

        self.advisor_unassigned = User.objects.create_user(
            email='advisor2@test.edu',
            password='password123',
            role='ASESOR',
            first_name='Mario',
            last_name='Externo'
        )

        self.student_user_1 = User.objects.create_user(
            email='student1@test.edu',
            password='password123',
            role='ESTUDIANTE',
            first_name='Ana',
            last_name='García'
        )

        self.student_user_2 = User.objects.create_user(
            email='student2@test.edu',
            password='password123',
            role='ESTUDIANTE',
            first_name='Pedro',
            last_name='López'
        )

        # 2. Crear Estudiantes
        self.student_1 = Student.objects.create(
            user=self.student_user_1,
            matricula='DOC-2025-001',
            nombre_completo='Ana García Gómez',
            programa_doctoral='Doctorado en Ciencias Computacionales',
            cohorte='2025-A',
            estatus_activo=True
        )

        self.student_2 = Student.objects.create(
            user=self.student_user_2,
            matricula='DOC-2025-002',
            nombre_completo='Pedro López Pérez',
            programa_doctoral='Doctorado en Ciencias Computacionales',
            cohorte='2025-A',
            estatus_activo=True
        )

        # 3. Asignar Semestres
        self.sem_1 = Semester.objects.create(
            student=self.student_1,
            numero=1,
            fecha_inicio=timezone.now().date() - timezone.timedelta(days=180),
            fecha_fin=timezone.now().date() - timezone.timedelta(days=30),
            is_active=False
        )
        self.sem_2 = Semester.objects.create(
            student=self.student_1,
            numero=2,
            fecha_inicio=timezone.now().date() - timezone.timedelta(days=29),
            fecha_fin=timezone.now().date() + timezone.timedelta(days=120),
            is_active=True
        )

        # 4. Asignar Comité Tutorial
        self.comm_1 = AcademicCommittee.objects.create(
            student=self.student_1,
            user=self.advisor_assigned,
            rol_comite='ASESOR_PRINCIPAL',
            is_active=True
        )

        # 5. Crear Sesión de Tutoría, Participantes y Observaciones
        self.session = TutoringSession.objects.create(
            student=self.student_1,
            semester=self.sem_2,
            fecha_sesion=timezone.now().date() - timezone.timedelta(days=10),
            modalidad='PRESENCIAL',
            resumen='Revisión de avances del capítulo 2 de tesis.',
            proxima_reunion_fecha=timezone.now().date() + timezone.timedelta(days=20),
            proxima_reunion_notas='Presentar resultados preliminares.',
            created_by=self.advisor_assigned
        )
        TutoringParticipant.objects.create(
            session=self.session,
            user=self.student_user_1,
            rol_en_sesion='ESTUDIANTE',
            asistencia=True
        )
        TutoringParticipant.objects.create(
            session=self.session,
            user=self.advisor_assigned,
            rol_en_sesion='ASESOR_PRINCIPAL',
            asistencia=True
        )
        TutoringObservation.objects.create(
            session=self.session,
            autor=self.advisor_assigned,
            titulo_tema='Metodología y Análisis',
            contenido='Se validaron las pruebas estadísticas realizadas.'
        )

        # 6. Crear Acuerdos
        self.agr_1 = Agreement.objects.create(
            student=self.student_1,
            session=self.session,
            descripcion='Completar marco experimental de la propuesta.',
            responsable=self.student_user_1,
            fecha_limite=timezone.now().date() + timezone.timedelta(days=15),
            estado=Agreement.STATUS_EN_PROCESO,
            created_by=self.advisor_assigned
        )
        AgreementAuditLog.objects.create(
            agreement=self.agr_1,
            user=self.advisor_assigned,
            estado_anterior='PENDIENTE',
            estado_nuevo='EN_PROCESO',
            comentario='Iniciado formalmente.'
        )

        # 7. Crear Avance de Tesis
        self.thesis = ThesisProgress.objects.create(
            student=self.student_1,
            semester=self.sem_2,
            porcentaje_avance=65,
            componentes_json={
                'protocolo': 100,
                'estadoArte': 90,
                'marcoTeorico': 80,
                'metodologia': 70,
                'analisis': 40,
                'redaccion': 30
            },
            observaciones='Excelente avance general.',
            fecha_registro=timezone.now().date()
        )

        # 8. Crear Evidencias
        self.evidence = Evidence.objects.create(
            student=self.student_1,
            semester=self.sem_2,
            tipo='ENLACE_DOI',
            actividad_tipo='OTRO',
            titulo='DOI Artículo Scopus',
            enlace_url='https://doi.org/10.1016/j.nexus.2025.01.001',
            fecha_carga=timezone.now().date(),
            created_by=self.student_user_1
        )

        # 9. Crear Producción Académica
        self.pub = Publication.objects.create(
            student=self.student_1,
            semester=self.sem_2,
            titulo='Deep Learning for Smart Universities',
            autores_texto='García, A., Asesor, C.',
            tipo='ARTICULO_JCR',
            revista_editorial='IEEE Transactions on Education',
            estado='PUBLICADO',
            fecha_publicacion=timezone.now().date() - timezone.timedelta(days=15),
            doi_url='https://doi.org/10.1016/j.nexus.2025.01.001',
            evidencia=self.evidence
        )

        self.event = AcademicEvent.objects.create(
            student=self.student_1,
            semester=self.sem_2,
            tipo_evento='CONGRESO_INTERNACIONAL',
            nombre_evento='IEEE International Conference on Higher Ed',
            titulo_ponencia='Orchestrating AI for Doctoral Tracking',
            fecha_presentacion=timezone.now().date() - timezone.timedelta(days=5),
            sede_lugar='San Francisco, CA',
            modalidad='PRESENCIAL'
        )

        self.stay = ResearchStay.objects.create(
            student=self.student_1,
            institucion_receptora='MIT Media Lab',
            pais='Estados Unidos',
            fecha_inicio=timezone.now().date() - timezone.timedelta(days=60),
            fecha_fin=timezone.now().date() - timezone.timedelta(days=30),
            responsable_estancia='Dr. John Doe',
            objetivos='Desarrollo de modelos neuronales.',
            resultados='Prototipo funcional evaluado.'
        )

        self.other = OtherProduct.objects.create(
            student=self.student_1,
            tipo_producto='SOFTWARE',
            titulo='NEXUS Pipeline Framework',
            descripcion='Librería Python para análisis de datos doctorales.',
            fecha_registro=timezone.now().date()
        )

    def test_export_student_excel_success(self):
        """Validar exportación XLSX del expediente de un estudiante (HTTP 200, 6 hojas, contenido correcto)."""
        self.client.force_authenticate(user=self.coordinator)
        url = f"/api/v2/reporting/students/{self.student_1.id}/export/?format=xlsx"
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('Expediente_DOC-2025-001.xlsx', response['Content-Disposition'])

        # Cargar libro con openpyxl
        excel_file = io.BytesIO(response.content)
        wb = openpyxl.load_workbook(excel_file)

        # Validar las 6 hojas requeridas
        expected_sheets = [
            "Datos Generales",
            "Tutorías",
            "Acuerdos",
            "Avance de Tesis",
            "Producción Académica",
            "Evidencias y DOIs"
        ]
        self.assertEqual(wb.sheetnames, expected_sheets)

        # Validar contenido en Hoja 1
        ws1 = wb["Datos Generales"]
        self.assertIn("DOC-2025-001", [str(cell.value) for row in ws1.rows for cell in row])
        self.assertIn("Ana García Gómez", [str(cell.value) for row in ws1.rows for cell in row])

        # Validar contenido en Hoja 2 (Tutorías)
        ws2 = wb["Tutorías"]
        self.assertTrue(any("Revisión de avances" in str(cell.value) for row in ws2.rows for cell in row))

        # Validar contenido en Hoja 3 (Acuerdos)
        ws3 = wb["Acuerdos"]
        self.assertTrue(any("Completar marco experimental" in str(cell.value) for row in ws3.rows for cell in row))

        # Validar contenido en Hoja 4 (Avance de Tesis)
        ws4 = wb["Avance de Tesis"]
        self.assertTrue(any("65%" in str(cell.value) for row in ws4.rows for cell in row))

        # Validar contenido en Hoja 5 (Producción Académica)
        ws5 = wb["Producción Académica"]
        self.assertTrue(any("Deep Learning for Smart Universities" in str(cell.value) for row in ws5.rows for cell in row))

        # Validar contenido en Hoja 6 (Evidencias y DOIs)
        ws6 = wb["Evidencias y DOIs"]
        self.assertTrue(any("https://doi.org/10.1016/j.nexus.2025.01.001" in str(cell.value) for row in ws6.rows for cell in row))

    def test_export_student_pdf_success(self):
        """Validar exportación PDF del expediente de un estudiante (HTTP 200, MIME type application/pdf, binario válido)."""
        self.client.force_authenticate(user=self.coordinator)
        url = f"/api/v2/reporting/students/{self.student_1.id}/export/?format=pdf"
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('Expediente_DOC-2025-001.pdf', response['Content-Disposition'])

        # Verificar cabecera mágica de PDF
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertGreater(len(response.content), 1000)

    def test_export_default_format_is_xlsx(self):
        """Validar que sin parámetro 'format', el formato por defecto sea XLSX."""
        self.client.force_authenticate(user=self.coordinator)
        url = f"/api/v2/reporting/students/{self.student_1.id}/export/"
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    def test_export_invalid_format_returns_400(self):
        """Validar que un formato desconocido retorne 400 Bad Request."""
        self.client.force_authenticate(user=self.coordinator)
        url = f"/api/v2/reporting/students/{self.student_1.id}/export/?format=csv"
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('format', response.data)

    def test_export_nonexistent_student_returns_404(self):
        """Validar que un estudiante no existente retorne 404 Not Found."""
        self.client.force_authenticate(user=self.coordinator)
        url = "/api/v2/reporting/students/99999/export/?format=xlsx"
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_rbac_advisor_access_control(self):
        """Validar que el asesor asignado pueda exportar, pero el no asignado sea rechazado con 403."""
        # 1. Asesor asignado -> 200 OK
        self.client.force_authenticate(user=self.advisor_assigned)
        url = f"/api/v2/reporting/students/{self.student_1.id}/export/?format=xlsx"
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        # 2. Asesor no asignado -> 403 Forbidden
        self.client.force_authenticate(user=self.advisor_unassigned)
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_rbac_student_access_control(self):
        """Validar que el estudiante titular pueda exportar su expediente, pero no el de otros."""
        # 1. Estudiante titular sobre su propio expediente -> 200 OK
        self.client.force_authenticate(user=self.student_user_1)
        url_own = f"/api/v2/reporting/students/{self.student_1.id}/export/?format=pdf"
        response = self.client.get(url_own)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        # 2. Estudiante intentando exportar expediente de otro alumno -> 403 Forbidden
        url_other = f"/api/v2/reporting/students/{self.student_2.id}/export/?format=pdf"
        response = self.client.get(url_other)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_anonymous_user_returns_401(self):
        """Validar que un usuario no autenticado sea rechazado con 401 Unauthorized."""
        self.client.logout()
        url = f"/api/v2/reporting/students/{self.student_1.id}/export/?format=xlsx"
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_global_students_export_excel_success(self):
        """Validar exportación global de la cohorte de estudiantes en Excel."""
        self.client.force_authenticate(user=self.coordinator)
        url = "/api/v2/reporting/export-students/?format=xlsx"
        response = self.client.get(url)

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('Listado_Estudiantes_', response['Content-Disposition'])

        excel_file = io.BytesIO(response.content)
        wb = openpyxl.load_workbook(excel_file)
        ws = wb.active
        self.assertEqual(ws.title, "Resumen Cohorte")

        # Comprobar que los estudiantes aparecen en el archivo
        self.assertTrue(any("DOC-2025-001" in str(cell.value) for row in ws.rows for cell in row))
        self.assertTrue(any("DOC-2025-002" in str(cell.value) for row in ws.rows for cell in row))
