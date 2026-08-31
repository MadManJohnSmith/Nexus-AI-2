"""
Motor de exportación tabular multi-hoja en formato Microsoft Excel (.xlsx)
para el expediente del estudiante en el Sistema N.E.X.U.S. (HU-28).
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone


def generate_student_excel_dossier(student) -> bytes:
    """
    Genera un archivo Excel (.xlsx) estructurado con 6 hojas estilizadas:
    1. Datos Generales (Ficha alumno, cohorte, comité tutorial activo, semestres).
    2. Tutorías (Sesiones, fechas, modalidad, participantes, observaciones).
    3. Acuerdos (Compromisos, responsables, fechas límite, estados, bitácoras).
    4. Avance de Tesis (Semestres 1 a 6, porcentajes, desglose de componentes).
    5. Producción Académica (Publicaciones, Congresos, Estancias, Otros productos).
    6. Evidencias y DOIs (Enlaces persistentes y archivos).
    """
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    default_sheet = wb.active

    # Paleta institucional N.E.X.U.S.
    primary_color = "6365EF"
    dark_color = "2C1867"
    light_bg = "F5F7FB"
    border_color = "D0D5DD"

    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    section_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Calibri", size=10, color="334155")
    italic_font = Font(name="Calibri", size=9, italic=True, color="64748B")

    title_fill = PatternFill(start_color=dark_color, end_color=dark_color, fill_type="solid")
    section_fill = PatternFill(start_color="4E50DC", end_color="4E50DC", fill_type="solid")
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    zebra_fill = PatternFill(start_color=light_bg, end_color=light_bg, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color)
    )

    def apply_title_banner(ws, title_text, max_col=6):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        cell = ws.cell(row=1, column=1, value=title_text)
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.cell(row=2, column=1, value=f"N.E.X.U.S. - Expediente de {student.nombre_completo} ({student.matricula}) | Generado: {timezone.now().strftime('%Y-%m-%d %H:%M')}")
        ws.cell(row=2, column=1).font = italic_font
        ws.row_dimensions[2].height = 18

    def auto_fit_columns(ws, min_width=12, max_width=50):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Evitar celdas combinadas de la fila 1 para el cálculo del ancho
                if cell.row in [1, 2]:
                    continue
                val = str(cell.value or '')
                if '\n' in val:
                    lines = val.split('\n')
                    max_len = max(max_len, max(len(l) for l in lines))
                else:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(min(max_len + 4, max_width), min_width)

    # -------------------------------------------------------------
    # 1. HOJA 1: "Datos Generales"
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Datos Generales")
    apply_title_banner(ws1, "EXPEDIENTE DOCTORAL - FICHA DE DATOS GENERALES", max_col=5)

    current_row = 4
    # Ficha del Estudiante
    ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    sec_cell = ws1.cell(row=current_row, column=1, value="INFORMACIÓN GENERAL DEL ALUMNO")
    sec_cell.font = section_font
    sec_cell.fill = section_fill
    sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[current_row].height = 24
    current_row += 1

    info_data = [
        ("Matrícula:", student.matricula, "Estatus Activo:", "ACTIVO" if student.estatus_activo else "INACTIVO"),
        ("Nombre Completo:", student.nombre_completo, "Cohorte:", student.cohorte),
        ("Programa Doctoral:", student.programa_doctoral, "Fecha de Registro:", student.created_at.strftime('%Y-%m-%d %H:%M') if student.created_at else "N/A"),
        ("Correo Electrónico:", student.user.email if student.user else "No asignado", "Última Actualización:", student.updated_at.strftime('%Y-%m-%d %H:%M') if student.updated_at else "N/A"),
    ]

    for row_vals in info_data:
        ws1.row_dimensions[current_row].height = 20
        c1 = ws1.cell(row=current_row, column=1, value=row_vals[0])
        c1.font = bold_font
        c1.fill = zebra_fill
        c1.border = thin_border

        c2 = ws1.cell(row=current_row, column=2, value=row_vals[1])
        c2.font = regular_font
        c2.border = thin_border

        c3 = ws1.cell(row=current_row, column=3, value=row_vals[2])
        c3.font = bold_font
        c3.fill = zebra_fill
        c3.border = thin_border

        c4 = ws1.cell(row=current_row, column=4, value=row_vals[3])
        c4.font = regular_font
        c4.border = thin_border
        current_row += 1

    current_row += 1
    # Comité Tutorial Activo
    ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    sec_cell2 = ws1.cell(row=current_row, column=1, value="COMITÉ TUTORIAL ASIGNADO")
    sec_cell2.font = section_font
    sec_cell2.fill = section_fill
    sec_cell2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[current_row].height = 24
    current_row += 1

    committee_headers = ["Rol en Comité", "Nombre del Académico", "Correo Institucional", "Fecha Asignación", "Estatus"]
    for col_idx, h in enumerate(committee_headers, 1):
        c = ws1.cell(row=current_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws1.row_dimensions[current_row].height = 22
    current_row += 1

    committee_members = student.committee_members.select_related('user').all()
    if committee_members.exists():
        for idx, member in enumerate(committee_members):
            ws1.row_dimensions[current_row].height = 20
            row_fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                member.get_rol_comite_display(),
                member.user.get_full_name() or member.user.username,
                member.user.email,
                member.fecha_asignacion.strftime('%Y-%m-%d') if member.fecha_asignacion else "N/A",
                "ACTIVO" if member.is_active else "INACTIVO"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws1.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                if row_fill:
                    c.fill = row_fill
            current_row += 1
    else:
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        c = ws1.cell(row=current_row, column=1, value="No se han asignado miembros al comité tutorial.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        current_row += 1

    current_row += 1
    # Semestres Cursados
    ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    sec_cell3 = ws1.cell(row=current_row, column=1, value="PERIODOS / SEMESTRES REGISTRADOS")
    sec_cell3.font = section_font
    sec_cell3.fill = section_fill
    sec_cell3.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[current_row].height = 24
    current_row += 1

    sem_headers = ["Semestre", "Fecha Inicio", "Fecha Fin", "Estatus Periodo"]
    for col_idx, h in enumerate(sem_headers, 1):
        c = ws1.cell(row=current_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws1.row_dimensions[current_row].height = 22
    current_row += 1

    semesters = student.semesters.all().order_by('numero')
    if semesters.exists():
        for idx, sem in enumerate(semesters):
            ws1.row_dimensions[current_row].height = 20
            row_fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                f"Semestre {sem.numero}",
                sem.fecha_inicio.strftime('%Y-%m-%d') if sem.fecha_inicio else "N/A",
                sem.fecha_fin.strftime('%Y-%m-%d') if sem.fecha_fin else "N/A",
                "ACTIVO (EN CURSO)" if sem.is_active else "CONCLUIDO"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws1.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                if row_fill:
                    c.fill = row_fill
            current_row += 1
    else:
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        c = ws1.cell(row=current_row, column=1, value="No hay registros de semestres para este estudiante.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        current_row += 1

    auto_fit_columns(ws1)

    # -------------------------------------------------------------
    # 2. HOJA 2: "Tutorías"
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Tutorías")
    apply_title_banner(ws2, "HISTORIAL DE SESIONES DE TUTORÍA Y SEGUIMIENTO", max_col=8)

    tutoring_headers = [
        "ID Sesión", "Fecha Sesión", "Semestre", "Modalidad",
        "Resumen de la Sesión", "Próxima Reunión", "Participantes", "Observaciones Registradas"
    ]
    for col_idx, h in enumerate(tutoring_headers, 1):
        c = ws2.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws2.row_dimensions[4].height = 26

    tutoring_sessions = student.tutoring_sessions.select_related('semester', 'created_by').prefetch_related('participants__user', 'observations__autor').all().order_by('-fecha_sesion')
    current_row = 5

    if tutoring_sessions.exists():
        for idx, session in enumerate(tutoring_sessions):
            row_fill = zebra_fill if idx % 2 == 1 else None

            # Formatear participantes
            parts_text = "\n".join([
                f"• {p.user.get_full_name() or p.user.username} ({p.get_rol_en_sesion_display()}): {'Asistió' if p.asistencia else 'Faltó'}"
                for p in session.participants.all()
            ]) or "Sin participantes registrados"

            # Formatear observaciones
            obs_text = "\n".join([
                f"[{o.titulo_tema}] {o.contenido} ({o.autor.get_full_name() or o.autor.username})"
                for o in session.observations.all()
            ]) or "Sin observaciones registradas"

            proxima_info = session.proxima_reunion_fecha.strftime('%Y-%m-%d') if session.proxima_reunion_fecha else "No programada"
            if session.proxima_reunion_notas:
                proxima_info += f"\nNotas: {session.proxima_reunion_notas}"

            vals = [
                f"SES-{session.id:04d}",
                session.fecha_sesion.strftime('%Y-%m-%d'),
                f"Semestre {session.semester.numero}" if session.semester else "N/A",
                session.get_modalidad_display(),
                session.resumen,
                proxima_info,
                parts_text,
                obs_text
            ]

            for col_idx, val in enumerate(vals, 1):
                c = ws2.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill:
                    c.fill = row_fill

            ws2.row_dimensions[current_row].height = max(30, 18 * max(parts_text.count('\n') + 1, obs_text.count('\n') + 1, vals[4].count('\n') + 1))
            current_row += 1
    else:
        ws2.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)
        c = ws2.cell(row=5, column=1, value="No se han registrado sesiones de tutoría.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        ws2.row_dimensions[5].height = 24

    auto_fit_columns(ws2)

    # -------------------------------------------------------------
    # 3. HOJA 3: "Acuerdos"
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Acuerdos")
    apply_title_banner(ws3, "ACUERDOS Y COMPROMISOS ACADÉMICOS", max_col=8)

    agreement_headers = [
        "ID Acuerdo", "Sesión Vinculada", "Descripción del Compromiso",
        "Responsable Asignado", "Fecha Límite", "Estado Actual", "Fecha Conclusión", "Bitácora de Cambios"
    ]
    for col_idx, h in enumerate(agreement_headers, 1):
        c = ws3.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws3.row_dimensions[4].height = 26

    agreements = student.agreements.select_related('responsable', 'session', 'created_by').prefetch_related('audit_logs__user').all().order_by('-fecha_limite')
    current_row = 5

    # Estilos de color para estados de acuerdos
    state_fills = {
        'PENDIENTE': PatternFill(start_color="F6FCFE", end_color="F6FCFE", fill_type="solid"),
        'EN_PROCESO': PatternFill(start_color="FEF8F3", end_color="FEF8F3", fill_type="solid"),
        'CONCLUIDO': PatternFill(start_color="E9FEF1", end_color="E9FEF1", fill_type="solid"),
        'VENCIDO': PatternFill(start_color="F8F1FF", end_color="F8F1FF", fill_type="solid"),
    }
    state_fonts = {
        'PENDIENTE': Font(name="Calibri", size=10, bold=True, color="57949D"),
        'EN_PROCESO': Font(name="Calibri", size=10, bold=True, color="B57136"),
        'CONCLUIDO': Font(name="Calibri", size=10, bold=True, color="437E5C"),
        'VENCIDO': Font(name="Calibri", size=10, bold=True, color="A14D98"),
    }

    if agreements.exists():
        for idx, agr in enumerate(agreements):
            row_fill = zebra_fill if idx % 2 == 1 else None

            logs_text = "\n".join([
                f"[{log.fecha_cambio.strftime('%Y-%m-%d %H:%M')}] {log.estado_anterior} -> {log.estado_nuevo}" +
                (f" ({log.comentario})" if log.comentario else "") +
                (f" por {log.user.get_full_name() or log.user.username}" if log.user else "")
                for log in agr.audit_logs.all()
            ]) or "Sin bitácora registrada"

            vals = [
                f"ACU-{agr.id:04d}",
                f"SES-{agr.session_id:04d}" if agr.session_id else "Independiente",
                agr.descripcion,
                agr.responsable.get_full_name() or agr.responsable.username,
                agr.fecha_limite.strftime('%Y-%m-%d') if agr.fecha_limite else "N/A",
                agr.get_estado_display(),
                agr.fecha_conclusion.strftime('%Y-%m-%d') if agr.fecha_conclusion else "Pendiente",
                logs_text
            ]

            for col_idx, val in enumerate(vals, 1):
                c = ws3.cell(row=current_row, column=col_idx, value=val)
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)

                if col_idx == 6:  # Columna de estado con estilo píldora
                    c.fill = state_fills.get(agr.estado, row_fill)
                    c.font = state_fonts.get(agr.estado, bold_font)
                    c.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    c.font = regular_font
                    if row_fill:
                        c.fill = row_fill

            ws3.row_dimensions[current_row].height = max(26, 16 * max(logs_text.count('\n') + 1, agr.descripcion.count('\n') + 1))
            current_row += 1
    else:
        ws3.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)
        c = ws3.cell(row=5, column=1, value="No se han registrado acuerdos o compromisos.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        ws3.row_dimensions[5].height = 24

    auto_fit_columns(ws3)

    # -------------------------------------------------------------
    # 4. HOJA 4: "Avance de Tesis"
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Avance de Tesis")
    apply_title_banner(ws4, "SEGUIMIENTO Y EVOLUCIÓN DE TESIS DOCTORAL", max_col=10)

    thesis_headers = [
        "Semestre", "% Avance Global", "Fecha Registro", "Protocolo (%)",
        "Estado del Arte (%)", "Marco Teórico (%)", "Metodología (%)",
        "Análisis (%)", "Redacción (%)", "Observaciones del Asesor"
    ]
    for col_idx, h in enumerate(thesis_headers, 1):
        c = ws4.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws4.row_dimensions[4].height = 28

    thesis_entries = student.thesis_progresses.select_related('semester').all().order_by('semester__numero', '-fecha_registro')
    current_row = 5

    if thesis_entries.exists():
        for idx, entry in enumerate(thesis_entries):
            row_fill = zebra_fill if idx % 2 == 1 else None
            comps = entry.componentes_json or {}

            vals = [
                f"Semestre {entry.semester.numero}" if entry.semester else "N/A",
                f"{entry.porcentaje_avance}%",
                entry.fecha_registro.strftime('%Y-%m-%d'),
                f"{comps.get('protocolo', 0)}%",
                f"{comps.get('estadoArte', 0)}%",
                f"{comps.get('marcoTeorico', 0)}%",
                f"{comps.get('metodologia', 0)}%",
                f"{comps.get('analisis', 0)}%",
                f"{comps.get('redaccion', 0)}%",
                entry.observaciones or "Sin observaciones registradas"
            ]

            for col_idx, val in enumerate(vals, 1):
                c = ws4.cell(row=current_row, column=col_idx, value=val)
                c.border = thin_border
                c.alignment = Alignment(vertical="center", horizontal="center" if col_idx <= 9 else "left", wrap_text=True)

                if col_idx == 2:
                    c.font = bold_font
                else:
                    c.font = regular_font

                if row_fill:
                    c.fill = row_fill

            ws4.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws4.merge_cells(start_row=5, start_column=1, end_row=5, end_column=10)
        c = ws4.cell(row=5, column=1, value="No se han registrado avances de tesis doctoral.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        ws4.row_dimensions[5].height = 24

    auto_fit_columns(ws4)

    # -------------------------------------------------------------
    # 5. HOJA 5: "Producción Académica"
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Producción Académica")
    apply_title_banner(ws5, "PRODUCCIÓN CIENTÍFICA, CONGRESOS, ESTANCIAS Y OTROS PRODUCTOS", max_col=7)

    current_row = 4

    # 5.1 Publicaciones Científicas
    ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
    sec_c = ws5.cell(row=current_row, column=1, value="1. PUBLICACIONES CIENTÍFICAS (HU-17)")
    sec_c.font = section_font
    sec_c.fill = section_fill
    sec_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[current_row].height = 24
    current_row += 1

    pub_headers = ["Tipo Publicación", "Título del Trabajo", "Autores", "Revista / Editorial", "Estado", "Fecha Publicación", "DOI / Enlace"]
    for col_idx, h in enumerate(pub_headers, 1):
        c = ws5.cell(row=current_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws5.row_dimensions[current_row].height = 22
    current_row += 1

    pubs = student.publications.all().order_by('-fecha_publicacion', '-created_at')
    if pubs.exists():
        for idx, pub in enumerate(pubs):
            row_fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                pub.get_tipo_display(),
                pub.titulo,
                pub.autores_texto,
                pub.revista_editorial,
                pub.get_estado_display(),
                pub.fecha_publicacion.strftime('%Y-%m-%d') if pub.fecha_publicacion else "N/A",
                pub.doi_url or "N/A"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws5.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill:
                    c.fill = row_fill
            ws5.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c = ws5.cell(row=current_row, column=1, value="No hay publicaciones científicas registradas.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        current_row += 1

    current_row += 1

    # 5.2 Eventos Académicos
    ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
    sec_c = ws5.cell(row=current_row, column=1, value="2. CONGRESOS, COLOQUIOS Y SIMPOSIOS (HU-18)")
    sec_c.font = section_font
    sec_c.fill = section_fill
    sec_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[current_row].height = 24
    current_row += 1

    event_headers = ["Tipo Evento", "Nombre del Evento", "Título Ponencia", "Fecha Presentación", "Sede / Lugar", "Modalidad", "Semestre"]
    for col_idx, h in enumerate(event_headers, 1):
        c = ws5.cell(row=current_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws5.row_dimensions[current_row].height = 22
    current_row += 1

    events = student.academic_events.select_related('semester').all().order_by('-fecha_presentacion')
    if events.exists():
        for idx, ev in enumerate(events):
            row_fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                ev.get_tipo_evento_display(),
                ev.nombre_evento,
                ev.titulo_ponencia,
                ev.fecha_presentacion.strftime('%Y-%m-%d') if ev.fecha_presentacion else "N/A",
                ev.sede_lugar,
                ev.get_modalidad_display(),
                f"Semestre {ev.semester.numero}" if ev.semester else "N/A"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws5.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill:
                    c.fill = row_fill
            ws5.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c = ws5.cell(row=current_row, column=1, value="No hay eventos académicos registrados.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        current_row += 1

    current_row += 1

    # 5.3 Estancias de Investigación
    ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
    sec_c = ws5.cell(row=current_row, column=1, value="3. ESTANCIAS DE INVESTIGACIÓN (HU-19)")
    sec_c.font = section_font
    sec_c.fill = section_fill
    sec_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[current_row].height = 24
    current_row += 1

    stay_headers = ["Institución Receptora", "País", "Periodo Inicio", "Periodo Fin", "Investigador Anfitrión", "Objetivos", "Resultados"]
    for col_idx, h in enumerate(stay_headers, 1):
        c = ws5.cell(row=current_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws5.row_dimensions[current_row].height = 22
    current_row += 1

    stays = student.research_stays.all().order_by('-fecha_inicio')
    if stays.exists():
        for idx, stay in enumerate(stays):
            row_fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                stay.institucion_receptora,
                stay.pais,
                stay.fecha_inicio.strftime('%Y-%m-%d') if stay.fecha_inicio else "N/A",
                stay.fecha_fin.strftime('%Y-%m-%d') if stay.fecha_fin else "N/A",
                stay.responsable_estancia,
                stay.objetivos,
                stay.resultados
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws5.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill:
                    c.fill = row_fill
            ws5.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c = ws5.cell(row=current_row, column=1, value="No hay estancias de investigación registradas.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        current_row += 1

    current_row += 1

    # 5.4 Otros Productos
    ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    sec_c = ws5.cell(row=current_row, column=1, value="4. OTROS PRODUCTOS (SOFTWARE, PATENTES, BASES DE DATOS) (HU-20)")
    sec_c.font = section_font
    sec_c.fill = section_fill
    sec_c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws5.row_dimensions[current_row].height = 24
    current_row += 1

    other_headers = ["Tipo de Producto", "Título del Producto", "Descripción / Características", "Fecha Registro", "ID Evidencia"]
    for col_idx, h in enumerate(other_headers, 1):
        c = ws5.cell(row=current_row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws5.row_dimensions[current_row].height = 22
    current_row += 1

    others = student.other_products.select_related('evidencia').all().order_by('-fecha_registro')
    if others.exists():
        for idx, oth in enumerate(others):
            row_fill = zebra_fill if idx % 2 == 1 else None
            vals = [
                oth.get_tipo_producto_display(),
                oth.titulo,
                oth.descripcion,
                oth.fecha_registro.strftime('%Y-%m-%d') if oth.fecha_registro else "N/A",
                f"EVI-{oth.evidencia_id:04d}" if oth.evidencia_id else "N/A"
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws5.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill:
                    c.fill = row_fill
            ws5.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws5.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        c = ws5.cell(row=current_row, column=1, value="No hay otros productos académicos registrados.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        current_row += 1

    auto_fit_columns(ws5)

    # -------------------------------------------------------------
    # 6. HOJA 6: "Evidencias y DOIs"
    # -------------------------------------------------------------
    ws6 = wb.create_sheet(title="Evidencias y DOIs")
    apply_title_banner(ws6, "REPOSITORIO DE EVIDENCIAS DOCUMENTALES Y ENLACES DOI (HU-21/22)", max_col=8)

    evidence_headers = [
        "ID Evidencia", "Tipo", "Actividad Vinculada", "Título de la Evidencia",
        "Descripción", "Archivo / DOI URL", "Tamaño (KB)", "Fecha Carga"
    ]
    for col_idx, h in enumerate(evidence_headers, 1):
        c = ws6.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border
    ws6.row_dimensions[4].height = 26

    evidences = student.evidences.select_related('semester', 'created_by').all().order_by('-fecha_carga')
    current_row = 5

    if evidences.exists():
        for idx, evi in enumerate(evidences):
            row_fill = zebra_fill if idx % 2 == 1 else None

            recurso_text = evi.enlace_url if evi.tipo == 'ENLACE_DOI' else (evi.archivo_adjunto.name if evi.archivo_adjunto else "Sin archivo")
            act_info = f"{evi.get_actividad_tipo_display()}" + (f" (ID: {evi.actividad_id})" if evi.actividad_id else "")
            size_kb = f"{(evi.file_size_bytes / 1024):.1f} KB" if evi.file_size_bytes else "-"

            vals = [
                f"EVI-{evi.id:04d}",
                evi.get_tipo_display(),
                act_info,
                evi.titulo,
                evi.descripcion or "-",
                recurso_text,
                size_kb,
                evi.fecha_carga.strftime('%Y-%m-%d') if evi.fecha_carga else "N/A"
            ]

            for col_idx, val in enumerate(vals, 1):
                c = ws6.cell(row=current_row, column=col_idx, value=val)
                c.font = regular_font
                c.border = thin_border
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if row_fill:
                    c.fill = row_fill

            ws6.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws6.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)
        c = ws6.cell(row=5, column=1, value="No se han cargado evidencias documentales ni enlaces DOI.")
        c.font = italic_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        ws6.row_dimensions[5].height = 24

    auto_fit_columns(ws6)

    # Eliminar la hoja inicial en blanco
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    # Guardar en memoria y retornar bytes
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_cohort_summary_excel(students_queryset) -> bytes:
    """
    Genera un archivo Excel global con el resumen general de la cohorte / listado de estudiantes.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen Cohorte"

    primary_color = "6365EF"
    dark_color = "2C1867"
    light_bg = "F5F7FB"
    border_color = "D0D5DD"

    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=10, color="334155")
    italic_font = Font(name="Calibri", size=9, italic=True, color="64748B")

    title_fill = PatternFill(start_color=dark_color, end_color=dark_color, fill_type="solid")
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    zebra_fill = PatternFill(start_color=light_bg, end_color=light_bg, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color)
    )

    ws.merge_cells("A1:H1")
    cell = ws["A1"]
    cell.value = "SISTEMA N.E.X.U.S. - REPORTE GENERAL DE COHORTE ESTUDIANTIL"
    cell.font = title_font
    cell.fill = title_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    ws["A2"] = f"Generado: {timezone.now().strftime('%Y-%m-%d %H:%M')} | Total Estudiantes: {students_queryset.count()}"
    ws["A2"].font = italic_font
    ws.row_dimensions[2].height = 18

    headers = [
        "Matrícula", "Nombre Completo", "Programa Doctoral", "Cohorte",
        "Estatus", "Asesor Principal", "% Avance Tesis", "Total Acuerdos"
    ]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws.row_dimensions[4].height = 26

    current_row = 5
    for idx, student in enumerate(students_queryset):
        row_fill = zebra_fill if idx % 2 == 1 else None

        # Asesor principal
        asesor_member = student.committee_members.filter(rol_comite='ASESOR_PRINCIPAL', is_active=True).select_related('user').first()
        asesor_name = (asesor_member.user.get_full_name() or asesor_member.user.username) if asesor_member else "Sin asignar"

        # Último avance
        last_thesis = student.thesis_progresses.order_by('-fecha_registro', '-id').first()
        avance_str = f"{last_thesis.porcentaje_avance}%" if last_thesis else "0%"

        total_acuerdos = student.agreements.count()

        vals = [
            student.matricula,
            student.nombre_completo,
            student.programa_doctoral,
            student.cohorte,
            "ACTIVO" if student.estatus_activo else "INACTIVO",
            asesor_name,
            avance_str,
            total_acuerdos
        ]

        for col_idx, val in enumerate(vals, 1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = Alignment(vertical="center", horizontal="center" if col_idx in [1, 4, 5, 7, 8] else "left")
            if row_fill:
                c.fill = row_fill

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2]:
                continue
            val = str(cell.value or '')
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 40), 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
